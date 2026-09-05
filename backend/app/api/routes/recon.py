from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.api.serializers import match_out
from app.core.db import get_db
from app.core.enums import (
    AuditAction,
    DocumentType,
    InvoiceSource,
    MatchStatus,
    MismatchResolution,
    ReturnType,
)
from app.models import (
    Client,
    Document,
    DocumentVersion,
    Entity,
    InvoiceMatch,
    InvoiceRecord,
    ReconciliationRun,
    ReturnItem,
    TaxPeriod,
    User,
)
from app.schemas.requests import MismatchUpdate, ReconRunRequest
from app.services import audit, matching
from app.services.permissions import get_case_or_403, require_ca

router = APIRouter(prefix="/api", tags=["reconciliation"])


def _latest_version(db: Session, case_id: int, doc_type: DocumentType) -> Optional[DocumentVersion]:
    doc = db.execute(
        select(Document).where(Document.case_id == case_id, Document.doc_type == doc_type)
    ).scalars().first()
    if not doc or not doc.versions:
        return None
    return doc.versions[-1]


def _current_run(db: Session, case_id: int) -> Optional[ReconciliationRun]:
    return db.execute(
        select(ReconciliationRun)
        .where(ReconciliationRun.case_id == case_id, ReconciliationRun.is_superseded.is_(False))
        .order_by(ReconciliationRun.id.desc())
        .limit(1)
    ).scalars().first()


def _row_count(db: Session, case_id: int, source: InvoiceSource, version_id: Optional[int]) -> int:
    if not version_id:
        return 0
    return len(
        db.execute(
            select(InvoiceRecord.id).where(
                InvoiceRecord.case_id == case_id,
                InvoiceRecord.source == source,
                InvoiceRecord.document_version_id == version_id,
            )
        ).scalars().all()
    )


@router.get("/cases/{case_id}/recon")
def recon_status(
    case_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    case = get_case_or_403(db, user, case_id)
    gstr2b = _latest_version(db, case_id, DocumentType.GSTR2B)
    pr = _latest_version(db, case_id, DocumentType.PURCHASE_REGISTER)
    run = _current_run(db, case_id)
    return {
        "case_id": case_id,
        "gstr2b_version": (
            {
                "id": gstr2b.id,
                "version_no": gstr2b.version_no,
                "filename": gstr2b.original_filename,
                "rows": _row_count(db, case_id, InvoiceSource.GSTR2B, gstr2b.id),
            }
            if gstr2b
            else None
        ),
        "purchase_register_version": (
            {
                "id": pr.id,
                "version_no": pr.version_no,
                "filename": pr.original_filename,
                "rows": _row_count(db, case_id, InvoiceSource.PURCHASE_REGISTER, pr.id),
            }
            if pr
            else None
        ),
        "ready_to_run": bool(gstr2b and pr),
        "current_run": (
            {
                "id": run.id,
                "created_at": run.created_at,
                "params": run.params,
                "summary": run.summary,
                "stale": bool(
                    gstr2b and pr and (run.gstr2b_version_id != gstr2b.id or run.pr_version_id != pr.id)
                ),
            }
            if run
            else None
        ),
    }


@router.post("/cases/{case_id}/recon/run")
def run_recon(
    case_id: int,
    payload: ReconRunRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    require_ca(user)
    case = get_case_or_403(db, user, case_id)
    gstr2b = _latest_version(db, case_id, DocumentType.GSTR2B)
    pr = _latest_version(db, case_id, DocumentType.PURCHASE_REGISTER)
    if not gstr2b or not pr:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "Both GSTR-2B and the Purchase Register must be uploaded before reconciling",
        )

    item = db.execute(
        select(ReturnItem).where(
            ReturnItem.case_id == case_id, ReturnItem.return_type == ReturnType.PR_RECON
        )
    ).scalars().first()

    run = matching.run_reconciliation(
        db, user, case, item, pr.id, gstr2b.id, payload.amount_tolerance, payload.date_tolerance_days
    )
    db.commit()
    return {"run_id": run.id, "summary": run.summary, "params": run.params}


@router.get("/cases/{case_id}/recon/matches")
def list_matches(
    case_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
    match_status: Optional[MatchStatus] = None,
    resolution_status: Optional[MismatchResolution] = None,
    action_required: bool = False,
):
    get_case_or_403(db, user, case_id)
    run = _current_run(db, case_id)
    if not run:
        return {"run_id": None, "summary": None, "items": []}

    stmt = select(InvoiceMatch).where(InvoiceMatch.run_id == run.id)
    if match_status:
        stmt = stmt.where(InvoiceMatch.match_status == match_status)
    if resolution_status:
        stmt = stmt.where(InvoiceMatch.resolution_status == resolution_status)
    if action_required:
        stmt = stmt.where(InvoiceMatch.match_status != MatchStatus.EXACT_MATCH)
    rows = db.execute(stmt.order_by(InvoiceMatch.match_status, InvoiceMatch.id)).scalars().all()
    return {"run_id": run.id, "summary": run.summary, "items": [match_out(m) for m in rows]}


@router.get("/cases/{case_id}/recon/reports")
def recon_reports(
    case_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)
):
    """The three Stage-1 reports: matched, missing in GSTR-2B, missing in PR --
    plus the finer categories the engine already produces."""
    get_case_or_403(db, user, case_id)
    run = _current_run(db, case_id)
    if not run:
        return {"run_id": None, "summary": None, "reports": {}}

    rows = db.execute(
        select(InvoiceMatch).where(InvoiceMatch.run_id == run.id)
    ).scalars().all()

    def bucket(*statuses):
        return [match_out(m) for m in rows if MatchStatus(m.match_status) in statuses]

    return {
        "run_id": run.id,
        "summary": run.summary,
        "reports": {
            "matched": bucket(MatchStatus.EXACT_MATCH),
            "missing_in_gstr2b": bucket(MatchStatus.MISSING_IN_2B),
            "missing_in_purchase_register": bucket(MatchStatus.MISSING_IN_PR),
            "partial_match": bucket(MatchStatus.PARTIAL_MATCH),
            "probable_match": bucket(MatchStatus.PROBABLE_MATCH),
            "mismatch": bucket(MatchStatus.MISMATCH),
        },
    }


@router.patch("/recon/matches/{match_id}")
def update_match(
    match_id: int,
    payload: MismatchUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    match = db.get(InvoiceMatch, match_id)
    if not match:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Match not found")
    case = get_case_or_403(db, user, match.case_id)

    data = payload.model_dump(exclude_unset=True)
    for field, value in data.items():
        if value is not None:
            setattr(match, field, value)
    if data.get("resolution_status") == MismatchResolution.RESOLVED:
        match.resolved_at = datetime.utcnow()
        match.resolved_by_user_id = user.id

    audit.record(
        db, user, AuditAction.MISMATCH_UPDATED, "InvoiceMatch",
        f"Mismatch #{match.id} updated ({match.match_status})",
        target_id=match.id, client_id=case.client_id, case_id=case.id,
        meta=payload.model_dump(mode="json", exclude_unset=True),
    )
    db.commit()
    db.refresh(match)
    return match_out(match)


# The report gets sent to clients and to their suppliers' accountants, so each
# section can be exported on its own -- nobody should have to be handed the
# whole reconciliation to be shown the three invoices they are missing.
EXPORT_SECTIONS = {
    "matched": ("Matched", [MatchStatus.EXACT_MATCH]),
    "missing_in_gstr2b": ("Missing in GSTR-2B", [MatchStatus.MISSING_IN_2B]),
    "missing_in_purchase_register": ("Missing in PR", [MatchStatus.MISSING_IN_PR]),
    "mismatch": ("Mismatch", [MatchStatus.MISMATCH]),
    "partial_match": ("Partial match", [MatchStatus.PARTIAL_MATCH]),
    "probable_match": ("Probable match", [MatchStatus.PROBABLE_MATCH]),
}

EXPORT_HEADERS = [
    "Match Status", "Supplier GSTIN", "Supplier Name", "Invoice No", "Invoice Date",
    "PR Taxable", "PR IGST", "PR CGST", "PR SGST", "PR Cess",
    "2B Taxable", "2B IGST", "2B CGST", "2B SGST", "2B Cess",
    "Taxable Diff", "Tax Diff", "Diff Flags", "Resolution", "CA Remark", "Client Response",
]


def _match_row(m: InvoiceMatch) -> list:
    pr, tb = m.pr_record, m.gstr2b_record
    ref = pr or tb
    return [
        m.match_status if isinstance(m.match_status, str) else m.match_status.value,
        ref.supplier_gstin if ref else "",
        ref.supplier_name if ref else "",
        ref.invoice_no if ref else "",
        ref.invoice_date.isoformat() if ref and ref.invoice_date else "",
        pr.taxable_value if pr else "", pr.igst if pr else "",
        pr.cgst if pr else "", pr.sgst if pr else "", pr.cess if pr else "",
        tb.taxable_value if tb else "", tb.igst if tb else "",
        tb.cgst if tb else "", tb.sgst if tb else "", tb.cess if tb else "",
        m.taxable_value_diff, m.tax_diff, ", ".join(m.diff_flags or []),
        m.resolution_status if isinstance(m.resolution_status, str) else m.resolution_status.value,
        m.ca_remark or "", m.client_response or "",
    ]


def _write_sheet(sheet, rows: list) -> None:
    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for m in rows:
        sheet.append(_match_row(m))
    sheet.freeze_panes = "A2"
    # Enough width to read without dragging every column open.
    for i, header in enumerate(EXPORT_HEADERS, start=1):
        longest = max([len(header)] + [len(str(r[i - 1] or "")) for r in map(_match_row, rows)] or [0])
        sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 42)


def _cover_sheet(sheet, case, period, run, entity, client, section_label: str) -> None:
    """Who this is about and what was compared. A bare grid of invoice rows
    means nothing to whoever it gets forwarded to."""
    counts = (run.summary or {}).get("counts", {})
    params = run.params or {}
    rows = [
        ("GSTR-2B / Purchase Register reconciliation", ""),
        ("", ""),
        ("Client", client.name if client else ""),
        ("File number", entity.file_number if entity else ""),
        ("GSTIN", entity.gstin if entity else ""),
        ("Period", period.label if period else ""),
        ("Section", section_label),
        ("Reconciled on", run.created_at.strftime("%d %b %Y %H:%M") if run.created_at else ""),
        ("", ""),
        ("Purchase Register rows", params.get("pr_rows", "")),
        ("GSTR-2B rows", params.get("gstr2b_rows", "")),
        ("Amount tolerance", params.get("amount_tolerance", "")),
        ("Date tolerance (days)", params.get("date_tolerance_days", "")),
        ("", ""),
        ("Exact matches", counts.get(MatchStatus.EXACT_MATCH.value, 0)),
        ("Partial matches", counts.get(MatchStatus.PARTIAL_MATCH.value, 0)),
        ("Probable matches", counts.get(MatchStatus.PROBABLE_MATCH.value, 0)),
        ("Mismatches", counts.get(MatchStatus.MISMATCH.value, 0)),
        ("Missing in GSTR-2B", counts.get(MatchStatus.MISSING_IN_2B.value, 0)),
        ("Missing in Purchase Register", counts.get(MatchStatus.MISSING_IN_PR.value, 0)),
        ("", ""),
        ("Total lines compared", (run.summary or {}).get("total", 0)),
        ("Exact match rate", f"{(run.summary or {}).get('match_rate', 0)}%"),
    ]
    for label, value in rows:
        sheet.append([label, value])
    sheet["A1"].font = Font(bold=True, size=13)
    for row in sheet.iter_rows(min_row=3, min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 40


@router.get("/cases/{case_id}/recon/export")
def export_recon(
    case_id: int,
    section: str = "all",
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """The whole reconciliation, or one section of it, as a workbook to send on.

    section = all, or any of the report keys the UI shows as tabs.
    """
    case = get_case_or_403(db, user, case_id)
    run = _current_run(db, case_id)
    if not run:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "No reconciliation run for this case")
    if section != "all" and section not in EXPORT_SECTIONS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unknown section. Use all, " + ", ".join(EXPORT_SECTIONS),
        )

    rows = db.execute(select(InvoiceMatch).where(InvoiceMatch.run_id == run.id)).scalars().all()
    period = db.get(TaxPeriod, case.tax_period_id)
    entity = db.get(Entity, case.entity_id)
    client = db.get(Client, case.client_id)

    wanted = EXPORT_SECTIONS if section == "all" else {section: EXPORT_SECTIONS[section]}
    label = "All sections" if section == "all" else EXPORT_SECTIONS[section][0]

    workbook = Workbook()
    _cover_sheet(workbook.active, case, period, run, entity, client, label)
    workbook.active.title = "Summary"
    for name, statuses in wanted.values():
        _write_sheet(
            workbook.create_sheet(title=name),
            [m for m in rows if MatchStatus(m.match_status) in statuses],
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    stem = f"recon_{entity.gstin if entity else case_id}_{period.code if period else ''}"
    suffix = "" if section == "all" else "_" + section
    audit.record(
        db, user, AuditAction.DOWNLOAD, "ReconciliationRun",
        f"Reconciliation exported ({label})",
        target_id=run.id, client_id=case.client_id, case_id=case.id,
    )
    db.commit()
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{stem}{suffix}.xlsx"'},
    )
