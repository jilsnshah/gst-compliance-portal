"""Bulk import of the firm's existing client list.

    python import_clients.py "YRTURN_MORETHAN5CR_25-26 jils.xls"
    python import_clients.py clients.csv --dry-run

The sheet the firm keeps has four columns -- SR NO, FILE NO, NAME, GSTIN --
under a title row, with blank spacer rows between entries. Only those three
facts are known, so that is all this imports: one client, one file, one login
each -- plus the PAN, which is characters 3-12 of the GSTIN. Address and
contact details stay empty and are filled in later from Clients & files.

The file number is the key. Re-running skips anything already imported, so a
half-finished import can simply be run again.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from sqlalchemy import select

from app.core.db import SessionLocal
from app.core.enums import Constitution, FilingFrequency, Role
from app.core.security import hash_password
from app.models import Client, ClientUser, Entity, User

GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][0-9A-Z]Z[0-9A-Z]$")
PAN_RE = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]$")
WANTED = {"file no": "file_no", "name": "name", "gstin": "gstin"}

# The GSTIN's first two digits are the state. Only the ones the firm actually
# has files in -- an unknown code is left blank rather than guessed.
STATES = {
    "08": "Rajasthan", "23": "Madhya Pradesh", "24": "Gujarat",
    "27": "Maharashtra", "29": "Karnataka", "33": "Tamil Nadu",
    "06": "Haryana", "07": "Delhi", "09": "Uttar Pradesh", "19": "West Bengal",
}


def pan_from(gstin: str) -> str:
    pan = gstin[2:12]
    return pan if PAN_RE.match(pan) else ""


def read_rows(path: Path) -> list[dict]:
    """Finds the header row wherever it sits and reads the columns by name."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            grid = [[c.strip() for c in row] for row in csv.reader(fh)]
    elif path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            sys.exit("Reading .xls needs xlrd:  pip install xlrd")
        sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
        grid = [[str(c.value).strip() for c in sheet.row(r)] for r in range(sheet.nrows)]
    else:
        from openpyxl import load_workbook

        sheet = load_workbook(str(path), read_only=True, data_only=True).worksheets[0]
        grid = [["" if c is None else str(c).strip() for c in row]
                for row in sheet.iter_rows(values_only=True)]

    header_at = next(
        (i for i, row in enumerate(grid)
         if WANTED.keys() <= {c.lower() for c in row}),
        None,
    )
    if header_at is None:
        sys.exit(f"No header row with {', '.join(WANTED)} found in {path.name}")
    columns = {WANTED[c.lower()]: i for i, c in enumerate(grid[header_at]) if c.lower() in WANTED}

    rows = []
    for row in grid[header_at + 1:]:
        entry = {k: (row[i] if i < len(row) else "") for k, i in columns.items()}
        if not any(entry.values()):  # spacer row
            continue
        entry["gstin"] = entry["gstin"].upper().replace(" ", "")
        rows.append(entry)
    return rows


def import_rows(rows: list[dict], dry_run: bool = False) -> None:
    db = SessionLocal()
    created = skipped = rejected = 0
    try:
        for row in rows:
            file_no, name, gstin = row["file_no"], row["name"], row["gstin"]
            if not (file_no and name and gstin):
                print(f"  reject {file_no or '?':6} {name[:34]:34} missing a field")
                rejected += 1
                continue
            if not GSTIN_RE.match(gstin):
                print(f"  reject {file_no:6} {name[:34]:34} bad GSTIN {gstin}")
                rejected += 1
                continue

            existing = db.execute(
                select(Entity).where(Entity.file_number == file_no)
            ).scalars().first()
            if existing:
                skipped += 1
                continue
            if db.execute(select(Entity).where(Entity.gstin == gstin)).scalars().first():
                print(f"  reject {file_no:6} {name[:34]:34} GSTIN already on another file")
                rejected += 1
                continue

            # Stage 1: the client does not sign in, but the login exists so they
            # can be switched on without touching the data.
            email = f"{gstin.lower()}@gmail.com"
            if db.execute(select(User).where(User.email == email)).scalars().first():
                print(f"  reject {file_no:6} {name[:34]:34} {email} already has a login")
                rejected += 1
                continue

            client = Client(name=name)
            db.add(client)
            db.flush()

            login = User(
                email=email,
                full_name=name,
                hashed_password=hash_password(gstin),
                role=Role.CLIENT,
            )
            db.add(login)
            db.flush()
            db.add(ClientUser(user_id=login.id, client_id=client.id, is_primary_contact=True))

            db.add(Entity(
                client_id=client.id,
                file_number=file_no,
                legal_name=name,
                # Characters 3-12 of a GSTIN are the holder's PAN, so this is
                # read off the number rather than guessed. Anything that does
                # not look like a PAN is left blank instead of stored wrong.
                pan=pan_from(gstin),
                gstin=gstin,
                state_code=gstin[:2],
                state_name=STATES.get(gstin[:2]),
                constitution=Constitution.OTHER,
                filing_frequency=FilingFrequency.MONTHLY,
                applicable_services=[],
            ))
            created += 1

        if dry_run:
            db.rollback()
            print(f"\ndry run: would create {created}, skip {skipped}, reject {rejected}")
        else:
            db.commit()
            print(f"\ncreated {created}, already present {skipped}, rejected {rejected}")
    finally:
        db.close()


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if not args:
        sys.exit(__doc__)
    path = Path(args[0])
    if not path.exists():
        sys.exit(f"No such file: {path}")
    rows = read_rows(path)
    print(f"{len(rows)} entries in {path.name}")
    import_rows(rows, dry_run="--dry-run" in sys.argv)
